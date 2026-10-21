import json
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor


def add_policy_content(document, markdown):
    lines = [line.strip() for line in markdown.splitlines() if line.strip()]
    for index, line in enumerate(lines):
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            paragraph = document.add_paragraph(line[3:], style="Heading 2")
            paragraph.paragraph_format.keep_with_next = True
        elif line.startswith("### "):
            paragraph = document.add_paragraph(line[4:], style="Heading 3")
            paragraph.paragraph_format.keep_with_next = True
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            paragraph = document.add_paragraph(line)
            if index + 1 < len(lines) and lines[index + 1].startswith("- "):
                paragraph.paragraph_format.keep_with_next = True


def main():
    with open(sys.argv[1], encoding="utf-8") as source:
        policy = json.load(source)

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(28, 35, 33)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in (
        ("Heading 1", 19, 0, 8),
        ("Heading 2", 12, 16, 6),
        ("Heading 3", 10.5, 12, 5),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(16, 41, 26)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = document.styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet.font.size = Pt(10)
    bullet.paragraph_format.space_after = Pt(4)

    document.add_paragraph(policy["title"].upper(), style="Heading 1")
    effective = document.add_paragraph(f'Effective Date: {policy["effectiveDate"]}')
    effective.runs[0].font.size = Pt(9)
    effective.runs[0].font.color.rgb = RGBColor(83, 97, 91)
    effective.paragraph_format.space_after = Pt(18)
    add_policy_content(document, policy["body"])
    document.save(sys.argv[2])


if __name__ == "__main__":
    main()
